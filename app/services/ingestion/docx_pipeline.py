import asyncio
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from typing import TypedDict

import structlog
from docx import Document

from app.config import settings
from app.prompts import DOCUMENT_CLASSIFICATION_PROMPT
from app.services.llm.client import invoke_with_text, retry_llm_call

logger = structlog.get_logger(__name__)
_executor = ThreadPoolExecutor()

FIRST_PAGE_PARAGRAPH_COUNT = 8


class DocxResult(TypedDict):
    text: str
    type: str | None  # "exhibit", "filed_doc", "ic_notes", or None if classification is disabled


def _extract_text_sync(docx_bytes: bytes) -> tuple[str, str]:
    """Returns (first_page_text, full_text). first_page_text is the first
    FIRST_PAGE_PARAGRAPH_COUNT non-empty paragraphs, standing in for a PDF's
    first-page OCR text — .docx has no fixed pagination without rendering
    the document, so paragraph count is the closest equivalent sample."""
    document = Document(BytesIO(docx_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))
    full_text = "\n\n".join(paragraphs + table_lines)
    first_page_text = "\n\n".join(paragraphs[:FIRST_PAGE_PARAGRAPH_COUNT])
    return first_page_text, full_text


async def _no_classification() -> None:
    return None


async def _process_single(
    semaphore: asyncio.Semaphore,
    filename: str,
    docx_bytes: bytes,
) -> tuple[str, DocxResult | None, str | None]:
    async with semaphore:
        loop = asyncio.get_event_loop()
        logger.info("docx_processing_started", filename=filename, size_bytes=len(docx_bytes))
        try:
            first_page_text, full_text = await loop.run_in_executor(
                _executor, _extract_text_sync, docx_bytes
            )
            logger.info("docx_text_extracted", filename=filename, text_length=len(full_text))

            doc_type = (
                await loop.run_in_executor(
                    _executor,
                    retry_llm_call,
                    invoke_with_text,
                    DOCUMENT_CLASSIFICATION_PROMPT,
                    first_page_text,
                )
                if settings.CLASSIFY_PDF_LLM
                else await _no_classification()
            )
            if doc_type is not None:
                doc_type = doc_type.lower().strip()
                if doc_type not in ("exhibit", "filed_doc", "ic_notes"):
                    doc_type = "exhibit"
                logger.info("docx_classification_done", filename=filename, doc_type=doc_type)
            else:
                logger.info("docx_classification_skipped", filename=filename)

            return filename, DocxResult(text=full_text, type=doc_type), None
        except Exception as exc:
            logger.exception("docx_processing_failed", filename=filename, error=str(exc))
            return filename, None, str(exc)


async def run_docx_pipeline(
    files: dict[str, bytes],
) -> tuple[dict[str, DocxResult], dict[str, str]]:
    logger.info(
        "docx_pipeline_started",
        file_count=len(files),
        concurrency=settings.INGESTION_CONCURRENCY,
        classify_enabled=settings.CLASSIFY_PDF_LLM,
    )
    semaphore = asyncio.Semaphore(settings.INGESTION_CONCURRENCY)
    tasks = [_process_single(semaphore, name, data) for name, data in files.items()]
    outcomes = await asyncio.gather(*tasks)

    results: dict[str, DocxResult] = {}
    failed: dict[str, str] = {}
    for filename, result, error in outcomes:
        if error:
            failed[filename] = error
        else:
            results[filename] = result

    logger.info("docx_pipeline_done", success_count=len(results), failed_count=len(failed))
    return results, failed
