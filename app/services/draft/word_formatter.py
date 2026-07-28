import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

ALIGNMENT_MAP = {
    "[CENTER]": WD_ALIGN_PARAGRAPH.CENTER,
    "[RIGHT]": WD_ALIGN_PARAGRAPH.RIGHT,
    "[JUSTIFY]": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_TABLE_SEPARATOR_CELL = re.compile(r"^:?-+:?$")


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _parse_markdown_table(block: str) -> list[list[str]] | None:
    """Return a markdown table's rows (header + body) if `block` is one, else None."""
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if len(lines) < 2:
        return None
    if not all(line.startswith("|") and line.endswith("|") for line in lines[:2]):
        return None

    separator_cells = _split_table_row(lines[1])
    if not separator_cells or not all(
        _TABLE_SEPARATOR_CELL.match(cell) for cell in separator_cells
    ):
        return None

    header = _split_table_row(lines[0])
    if len(header) != len(separator_cells):
        return None

    rows = [header]
    for line in lines[2:]:
        if not (line.startswith("|") and line.endswith("|")):
            return None
        row = _split_table_row(line)
        if len(row) != len(header):
            return None
        rows.append(row)
    return rows


_BRACKET_CHARS = (":", ")")


def _strip_alignment_label(block: str) -> str:
    block = block.strip()
    for label in ALIGNMENT_MAP:
        if block.startswith(label):
            return block[len(label):].strip()
    return block


def _is_bracket_only_block(block: str) -> bool:
    return _strip_alignment_label(block) in _BRACKET_CHARS


def _has_internal_bracket(block: str) -> bool:
    content = _strip_alignment_label(block)
    return len(content) > 1 and any(char in content[:-1] for char in _BRACKET_CHARS)


def _drop_redundant_bracket_separators(blocks: list[str]) -> list[str]:
    """A caption block's margin character (":" or ")") sometimes has to appear as
    its own line between two different structural groups (e.g. a label and the
    row beneath it) — that's the template's own convention and stays. But a
    bracket-only line between two consecutive multi-field rows of the same kind
    (rows that already carry a bracket mid-line, not just trailing — e.g. several
    stacked party lines) is a duplicated separator that was never meant to repeat
    between them. Only that second case is dropped."""
    non_blank = [i for i, b in enumerate(blocks) if b.strip()]
    drop = set()
    for pos, i in enumerate(non_blank):
        if not _is_bracket_only_block(blocks[i]):
            continue
        if pos == 0 or pos == len(non_blank) - 1:
            continue
        prev_block, next_block = blocks[non_blank[pos - 1]], blocks[non_blank[pos + 1]]
        if _has_internal_bracket(prev_block) and _has_internal_bracket(next_block):
            drop.add(i)
    return [b for i, b in enumerate(blocks) if i not in drop]


def draft_to_docx_bytes(draft_text: str) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    blocks = _drop_redundant_bracket_separators(draft_text.split("\n\n"))
    for block in blocks:
        block = block.strip()
        if not block:
            continue

        table_rows = _parse_markdown_table(block)
        if table_rows is not None:
            table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            table.style = "Table Grid"
            for row_idx, row in enumerate(table_rows):
                for col_idx, cell_text in enumerate(row):
                    table.cell(row_idx, col_idx).text = cell_text
            continue

        alignment = WD_ALIGN_PARAGRAPH.LEFT
        for label, value in ALIGNMENT_MAP.items():
            if block.startswith(label):
                alignment = value
                block = block[len(label):].strip()
                break

        paragraph = doc.add_paragraph(block)
        paragraph.alignment = alignment

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
