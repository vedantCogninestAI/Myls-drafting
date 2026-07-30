import asyncio
import io
from concurrent.futures import ThreadPoolExecutor

import structlog
from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.config import settings
from app.services.shared.pdf_conversion import pdf_bytes_to_docx_bytes

logger = structlog.get_logger(__name__)
_executor = ThreadPoolExecutor()


def _iter_block_items(parent):
    """Yield paragraphs and tables in true document order (python-docx's own
    `.paragraphs`/`.tables` are separate flat lists that lose interleaving)."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError(f"unsupported parent type: {type(parent)}")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _runs_to_markdown(paragraph: Paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        if not run.text:
            continue
        run_text = run.text
        if run.bold and run.italic:
            run_text = f"***{run_text}***"
        elif run.bold:
            run_text = f"**{run_text}**"
        elif run.italic:
            run_text = f"*{run_text}*"
        parts.append(run_text)
    return "".join(parts).strip()


def _paragraph_to_markdown(paragraph: Paragraph) -> str:
    if not paragraph.text.strip():
        return ""
    text = _runs_to_markdown(paragraph)

    style_name = (paragraph.style.name or "").lower()
    if style_name.startswith("heading"):
        try:
            level = int(style_name.replace("heading", "").strip())
        except ValueError:
            level = 1
        md = f"{'#' * max(1, min(level, 6))} {text}"
    elif style_name == "title":
        md = f"# {text}"
    elif "list bullet" in style_name:
        md = f"- {text}"
    elif "list number" in style_name:
        md = f"1. {text}"
    else:
        md = text

    alignment = paragraph.alignment
    if alignment is not None and alignment.name != "LEFT":
        md = f"[{alignment.name}] {md}"
    return md


def _table_to_markdown(table: Table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _docx_to_markdown(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    blocks = []
    blank_run = 0
    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            md = _paragraph_to_markdown(block)
        else:
            md = _table_to_markdown(block)

        if not md:
            # A single blank paragraph is the normal default separator
            # (already implicit in the "\n\n" join below) — only a run of
            # 2 or more is unusual enough to be worth surfacing. Tables
            # can also parse to "" (no rows) but never represent a
            # deliberate spacing gap, so they don't count toward this.
            if isinstance(block, Paragraph):
                blank_run += 1
            continue

        if blank_run >= 2:
            blocks.append(f"[BLANK LINES: {blank_run}]")
        blank_run = 0
        blocks.append(md)
    return "\n\n".join(blocks)


async def _process_single(
    semaphore: asyncio.Semaphore,
    filename: str,
    file_bytes: bytes,
) -> tuple[str, str | None, str | None]:
    async with semaphore:
        loop = asyncio.get_event_loop()
        logger.info("template_processing_started", filename=filename, size_bytes=len(file_bytes))
        try:
            docx_bytes = file_bytes
            if filename.lower().endswith(".pdf"):
                docx_bytes = await loop.run_in_executor(
                    _executor, pdf_bytes_to_docx_bytes, file_bytes
                )
                logger.info(
                    "template_pdf_converted", filename=filename, docx_size_bytes=len(docx_bytes)
                )

            markdown_text = await loop.run_in_executor(_executor, _docx_to_markdown, docx_bytes)
            logger.info("template_processing_done", filename=filename, text_length=len(markdown_text))
            return filename, markdown_text, None

        except Exception as exc:
            logger.exception("template_processing_failed", filename=filename, error=str(exc))
            return filename, None, str(exc)


async def run_template_pipeline(
    files: dict[str, bytes],
) -> tuple[dict[str, str], dict[str, str]]:
    logger.info(
        "template_pipeline_started",
        file_count=len(files),
        concurrency=settings.TEMPLATE_CONCURRENCY,
    )
    semaphore = asyncio.Semaphore(settings.TEMPLATE_CONCURRENCY)
    tasks = [_process_single(semaphore, name, data) for name, data in files.items()]
    outcomes = await asyncio.gather(*tasks)

    results: dict[str, str] = {}
    failed: dict[str, str] = {}
    for filename, text, error in outcomes:
        if error:
            failed[filename] = error
        else:
            results[filename] = text

    logger.info(
        "template_pipeline_done",
        success_count=len(results),
        failed_count=len(failed),
    )
    return results, failed
